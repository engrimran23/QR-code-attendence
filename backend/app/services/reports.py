"""
Report generation: PDF (ReportLab) and Excel (openpyxl).
Each function returns raw bytes ready to stream as a file download.
"""
import io
from datetime import date, timedelta
from typing import List, Dict, Any

# ── Excel ────────────────────────────────────────────────────────────────────

def _wb_styles():
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    header_font   = Font(bold=True, color="FFFFFF", size=11)
    header_fill   = PatternFill("solid", fgColor="1F4E79")
    center        = Alignment(horizontal="center", vertical="center")
    present_fill  = PatternFill("solid", fgColor="C6EFCE")
    absent_fill   = PatternFill("solid", fgColor="FFC7CE")
    late_fill     = PatternFill("solid", fgColor="FFEB9C")
    thin          = Side(style="thin", color="AAAAAA")
    border        = Border(left=thin, right=thin, top=thin, bottom=thin)
    return header_font, header_fill, center, present_fill, absent_fill, late_fill, border


def daily_report_excel(
    school_name: str,
    branch_name: str,
    report_date: date,
    rows: List[Dict],
    summary: Dict,
) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    header_font, header_fill, center, p_fill, a_fill, l_fill, border = _wb_styles()

    wb = Workbook()
    ws = wb.active
    ws.title = f"Daily {report_date}"

    # Title block
    ws.merge_cells("A1:H1")
    ws["A1"] = f"{school_name} — {branch_name}"
    ws["A1"].font = Font(bold=True, size=14)
    ws["A1"].alignment = center

    ws.merge_cells("A2:H2")
    ws["A2"] = f"Daily Attendance Report — {report_date.strftime('%d %B %Y')}"
    ws["A2"].font = Font(bold=True, size=11)
    ws["A2"].alignment = center

    # Summary row
    ws["A3"] = f"Total: {summary['total']}   Present: {summary['present']}   Late: {summary['late']}   Absent: {summary['absent']}   Attendance: {summary['pct']}%"
    ws.merge_cells("A3:H3")
    ws["A3"].font = Font(italic=True, size=10)

    # Headers
    headers = ["#", "Roll No", "Student Name", "Class", "Section", "Status", "Check-in Time", "Manual"]
    for col, h in enumerate(headers, 1):
        cell = ws.cell(row=5, column=col, value=h)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = center
        cell.border = border

    # Data rows
    status_fill = {"present": p_fill, "late": l_fill, "absent": a_fill, "leave": l_fill}
    for i, row in enumerate(rows, 1):
        data = [i, row["roll_number"], row["student_name"], row["class_name"],
                row.get("section") or "", row["status"].upper(),
                row["check_in_time"].strftime("%I:%M %p") if row.get("check_in_time") else "—",
                "Yes" if row.get("is_manual") else "No"]
        for col, val in enumerate(data, 1):
            cell = ws.cell(row=5 + i, column=col, value=val)
            cell.border = border
            cell.alignment = center
            if col == 6:
                cell.fill = status_fill.get(row["status"], p_fill)

    ws.column_dimensions["C"].width = 25
    ws.column_dimensions["B"].width = 12
    ws.column_dimensions["G"].width = 14

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf.read()


def monthly_report_excel(
    school_name: str,
    branch_name: str,
    year: int,
    month: int,
    class_name: str,
    section: str,
    working_days: int,
    students: List[Dict],
    date_headers: List[date],
) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    import calendar
    header_font, header_fill, center, p_fill, a_fill, l_fill, border = _wb_styles()

    wb = Workbook()
    ws = wb.active
    month_name = calendar.month_name[month]
    ws.title = f"{month_name[:3]} {year}"

    # Title
    ws.merge_cells(f"A1:{_col(4 + len(date_headers))}1")
    ws["A1"] = f"{school_name} | {branch_name} | Class {class_name} {section or ''}"
    ws["A1"].font = Font(bold=True, size=13)
    ws["A1"].alignment = center

    ws.merge_cells(f"A2:{_col(4 + len(date_headers))}2")
    ws["A2"] = f"Monthly Attendance — {month_name} {year}  (Working days: {working_days})"
    ws["A2"].font = Font(bold=True, size=11)
    ws["A2"].alignment = center

    # Column headers: Roll | Name | date1 | date2 … | Present | Late | Absent | %
    fixed = ["Roll", "Name"]
    trailing = ["Present", "Late", "Absent", "%"]
    all_headers = fixed + [d.strftime("%d") for d in date_headers] + trailing

    for col, h in enumerate(all_headers, 1):
        cell = ws.cell(row=4, column=col, value=h)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = center
        cell.border = border

    for i, stu in enumerate(students, 1):
        row_num = 4 + i
        ws.cell(row=row_num, column=1, value=stu["roll_number"]).border = border
        ws.cell(row=row_num, column=2, value=stu["student_name"]).border = border

        col = 3
        present_count = late_count = absent_count = 0
        for d in date_headers:
            status = stu["days"].get(str(d), "—")
            cell = ws.cell(row=row_num, column=col, value=status[:1].upper() if status != "—" else "—")
            cell.border = border
            cell.alignment = center
            if status == "present":   cell.fill = p_fill;  present_count += 1
            elif status == "late":    cell.fill = l_fill;  late_count += 1
            elif status == "absent":  cell.fill = a_fill;  absent_count += 1
            col += 1

        pct = round((present_count + late_count) / working_days * 100, 1) if working_days else 0
        for val in [present_count, late_count, absent_count, f"{pct}%"]:
            c = ws.cell(row=row_num, column=col, value=val)
            c.border = border; c.alignment = center
            col += 1

    ws.column_dimensions["B"].width = 24
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf.read()


def _col(n: int) -> str:
    """Convert column number to Excel letter (1→A, 27→AA)."""
    result = ""
    while n:
        n, r = divmod(n - 1, 26)
        result = chr(65 + r) + result
    return result


# ── PDF ──────────────────────────────────────────────────────────────────────

def _pdf_styles():
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.lib.enums import TA_CENTER
    from reportlab.lib import colors
    styles = getSampleStyleSheet()
    title_style  = ParagraphStyle("title",  parent=styles["Title"],  fontSize=16, spaceAfter=4)
    sub_style    = ParagraphStyle("sub",    parent=styles["Normal"], fontSize=10, spaceAfter=12, alignment=TA_CENTER)
    return title_style, sub_style


def daily_report_pdf(
    school_name: str,
    branch_name: str,
    report_date: date,
    rows: List[Dict],
    summary: Dict,
) -> bytes:
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib import colors
    from reportlab.lib.units import cm
    from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.lib.enums import TA_CENTER

    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=landscape(A4),
                            leftMargin=1.5*cm, rightMargin=1.5*cm,
                            topMargin=1.5*cm, bottomMargin=1.5*cm)
    styles = getSampleStyleSheet()
    story = []

    story.append(Paragraph(f"{school_name} — {branch_name}", styles["Title"]))
    story.append(Paragraph(
        f"Daily Attendance Report &nbsp;|&nbsp; {report_date.strftime('%d %B %Y')}",
        styles["Normal"]))
    story.append(Paragraph(
        f"Total: {summary['total']} &nbsp;|&nbsp; Present: {summary['present']} &nbsp;|&nbsp; "
        f"Late: {summary['late']} &nbsp;|&nbsp; Absent: {summary['absent']} &nbsp;|&nbsp; "
        f"Attendance: {summary['pct']}%",
        styles["Normal"]))
    story.append(Spacer(1, 0.4*cm))

    header = ["#", "Roll No", "Student Name", "Class", "Section", "Status", "Check-in", "Manual"]
    data = [header]
    for i, row in enumerate(rows, 1):
        data.append([
            str(i),
            row["roll_number"],
            row["student_name"],
            row["class_name"],
            row.get("section") or "",
            row["status"].upper(),
            row["check_in_time"].strftime("%I:%M %p") if row.get("check_in_time") else "—",
            "Yes" if row.get("is_manual") else "No",
        ])

    col_widths = [1*cm, 2.2*cm, 5.5*cm, 2.2*cm, 2*cm, 2.2*cm, 2.8*cm, 1.8*cm]
    tbl = Table(data, colWidths=col_widths, repeatRows=1)
    tbl.setStyle(TableStyle([
        ("BACKGROUND",  (0, 0), (-1, 0), colors.HexColor("#1F4E79")),
        ("TEXTCOLOR",   (0, 0), (-1, 0), colors.white),
        ("FONTNAME",    (0, 0), (-1, 0), "Helvetica-Bold"),
        ("FONTSIZE",    (0, 0), (-1, -1), 8),
        ("ALIGN",       (0, 0), (-1, -1), "CENTER"),
        ("GRID",        (0, 0), (-1, -1), 0.5, colors.grey),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F0F4F8")]),
    ]))

    # Colour status cells
    status_colors = {"PRESENT": "#C6EFCE", "LATE": "#FFEB9C", "ABSENT": "#FFC7CE", "LEAVE": "#FFEB9C"}
    for row_idx, row in enumerate(rows, 1):
        sc = status_colors.get(row["status"].upper())
        if sc:
            tbl.setStyle(TableStyle([("BACKGROUND", (5, row_idx), (5, row_idx), colors.HexColor(sc))]))

    story.append(tbl)
    doc.build(story)
    buf.seek(0)
    return buf.read()


def student_report_pdf(
    school_name: str,
    student_name: str,
    roll_number: str,
    class_name: str,
    from_date: date,
    to_date: date,
    records: List[Dict],
    stats: Dict,
) -> bytes:
    from reportlab.lib.pagesizes import A4
    from reportlab.lib import colors
    from reportlab.lib.units import cm
    from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
    from reportlab.lib.styles import getSampleStyleSheet

    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=A4,
                            leftMargin=2*cm, rightMargin=2*cm,
                            topMargin=2*cm, bottomMargin=2*cm)
    styles = getSampleStyleSheet()
    story = []

    story.append(Paragraph(school_name, styles["Title"]))
    story.append(Paragraph(f"Student Attendance Report", styles["h2"]))
    story.append(Spacer(1, 0.3*cm))

    info = [
        ["Student Name:", student_name, "Roll No:", roll_number],
        ["Class:", class_name, "Period:", f"{from_date} to {to_date}"],
        ["Total Days:", str(stats["total"]), "Present:", str(stats["present"])],
        ["Late:", str(stats["late"]), "Absent:", str(stats["absent"])],
        ["Attendance %:", f"{stats['pct']}%", "", ""],
    ]
    info_tbl = Table(info, colWidths=[3.5*cm, 5*cm, 3*cm, 5*cm])
    info_tbl.setStyle(TableStyle([
        ("FONTNAME", (0, 0), (0, -1), "Helvetica-Bold"),
        ("FONTNAME", (2, 0), (2, -1), "Helvetica-Bold"),
        ("FONTSIZE", (0, 0), (-1, -1), 9),
        ("GRID",     (0, 0), (-1, -1), 0.3, colors.lightgrey),
    ]))
    story.append(info_tbl)
    story.append(Spacer(1, 0.5*cm))

    data = [["#", "Date", "Day", "Status", "Check-in Time", "Note"]]
    for i, rec in enumerate(records, 1):
        data.append([
            str(i),
            str(rec["date"]),
            rec["date"].strftime("%A"),
            rec["status"].upper(),
            rec["check_in_time"].strftime("%I:%M %p") if rec.get("check_in_time") else "—",
            rec.get("note") or "",
        ])

    tbl = Table(data, colWidths=[1*cm, 2.8*cm, 2.8*cm, 2.5*cm, 3*cm, 4.5*cm], repeatRows=1)
    tbl.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1F4E79")),
        ("TEXTCOLOR",  (0, 0), (-1, 0), colors.white),
        ("FONTNAME",   (0, 0), (-1, 0), "Helvetica-Bold"),
        ("FONTSIZE",   (0, 0), (-1, -1), 8),
        ("ALIGN",      (0, 0), (-1, -1), "CENTER"),
        ("GRID",       (0, 0), (-1, -1), 0.5, colors.grey),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F0F4F8")]),
    ]))
    story.append(tbl)
    doc.build(story)
    buf.seek(0)
    return buf.read()
